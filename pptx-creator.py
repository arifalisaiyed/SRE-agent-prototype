from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Use Cases
slide_1 = prs.slides.add_slide(prs.slide_layouts[1])
title_1 = slide_1.shapes.title
title_1.text = "10 Use Cases for SRE Agent Prototyping"
content_1 = slide_1.placeholders[1]
content_1.text = (
    "1. Automated Incident Triage\n"
    "2. Context-Aware Remediation (Uptime/Diagnostics)\n"
    "3. Historical Pattern Matching (Cassandra lookup)\n"
    "4. OOM (Out of Memory) Resolution\n"
    "5. Disk Partition Management\n"
    "6. CPU Spike Correlation\n"
    "7. Database Connection Recovery\n"
    "8. Graceful Service Restart Synthesis\n"
    "9. Audit Logging & Post-Mortem Prep\n"
    "10. Predictive Alerting based on Kafka trends"
)

# Slide 2: Architecture
slide_2 = prs.slides.add_slide(prs.slide_layouts[1])
title_2 = slide_2.shapes.title
title_2.text = "High-Level Component Architecture"

# Creating a text-based block diagram for the slide
content_2 = slide_2.placeholders[1]
content_2.text = (
    "[Alert Source] --> [Kafka Topic] --> [SRE Agent]\n\n"
    "Inside SRE Agent:\n"
    "   - LangChain (Orchestrator)\n"
    "   - Groq/Llama 3.3 (Reasoning)\n"
    "   - Local Python (Tool Execution)\n\n"
    "Output Sink:\n"
    "   - Cassandra DB (Audit & Fix Logs)"
)

prs.save('SRE_Agent_Prototype.pptx')
print("PowerPoint saved as SRE_Agent_Prototype.pptx")